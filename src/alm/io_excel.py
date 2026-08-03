from __future__ import annotations

from pathlib import Path

from .aggregate import logic_map_from_edge_stats
from .models import Account, EdgeKey, EdgeStat, LogicMap
from .score import score_map


def load_excel_reference(xlsx_path: Path, *, window_label: str = "excel_ref") -> LogicMap:
    """Load pre-aggregated edges from the reference workbook (sheets rw, ca)."""
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required for build-from-excel; pip install openpyxl"
        ) from exc

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    accounts: dict[str, Account] = {}
    if "ca" in wb.sheetnames:
        ws = wb["ca"]
        for row in ws.iter_rows(min_row=1, values_only=True):
            if not row or row[0] is None:
                continue
            name = str(row[0]).strip()
            acct_type = str(row[1]).strip() if len(row) > 1 and row[1] is not None else "Unknown"
            accounts[name] = Account(account_id=name, name=name, account_type=acct_type)

    if "rw" not in wb.sheetnames:
        raise ValueError(f"workbook missing 'rw' sheet: {xlsx_path}")

    ws = wb["rw"]
    rows = ws.iter_rows(min_row=2, values_only=True)
    edges: dict[EdgeKey, EdgeStat] = {}
    for row in rows:
        if not row or row[1] is None or row[2] is None:
            continue
        debit = str(row[1]).strip()
        credit = str(row[2]).strip()
        weight = float(row[3] or 0)
        depth = int(float(row[4] or 0))
        if weight <= 0 or depth <= 0:
            continue
        key = EdgeKey(debit, credit)
        edges[key] = EdgeStat(key=key, weight_sum=weight, depth=depth)
        for aid in (debit, credit):
            if aid not in accounts:
                accounts[aid] = Account(account_id=aid, name=aid, account_type="Unknown")

    wb.close()
    logic_map = logic_map_from_edge_stats(
        accounts,
        edges,
        window_label=window_label,
        txn_count=0,
        line_count=0,
    )
    return score_map(logic_map)
